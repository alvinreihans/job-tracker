"""Excel "Job Tracker.xlsx" -> SQL insert buat tabel applications.

Pakai:  pip install openpyxl
        python scripts/import_xlsx.py "D:/Document/apply-for-a-job/Job Tracker.xlsx" email@kamu.com > supabase/seed.sql
Lalu jalankan supabase/seed.sql di Supabase SQL Editor. Aman dijalankan ulang (baris yang sama di-skip).
seed.sql berisi data pribadi dan ada di .gitignore. Jangan di-commit.
"""

import sys

import openpyxl

# Harus sama dengan lib/constants.ts dan check constraint di supabase/schema.sql.
STATUSES = {"Applied", "Seen", "Psikotes", "Interview HR", "Interview User", "Offering", "Diterima", "Ditolak"}
SOURCES = {"Jobstreet", "Kalibrr", "LinkedIn", "Astra Career", "Glints", "Website Perusahaan", "Referral", "Lainnya"}


def text(v):
    return str(v).strip() if v is not None and str(v).strip() else None


def sql(v):
    # standard_conforming_strings nyala di Supabase: cukup gandakan kutip tunggal.
    return "null" if v is None else "'" + v.replace("'", "''") + "'"


def read_details(ws):
    """Sheet Detail: baris dengan nama perusahaan = mulai lamaran baru, baris kosong = lanjutan poin."""
    details, key = {}, None
    for company, position, resp, req, *_ in ws.iter_rows(min_row=2, values_only=True):
        if text(company):
            key = (text(company), text(position))
            details[key] = ([], [])
        if key:
            if text(resp):
                details[key][0].append(text(resp))
            if text(req):
                details[key][1].append(text(req))
    return details


def main(path, email):
    wb = openpyxl.load_workbook(path, data_only=True)
    details = read_details(wb["Detail"])
    rows, problems = [], []

    for n, (company, position, applied, url, location, source, status, notes, *_) in enumerate(
        wb["Lamaran"].iter_rows(min_row=2, values_only=True), start=2
    ):
        company, position = text(company), text(position)
        if not company:
            continue
        source, status = text(source), text(status)
        if source not in SOURCES:
            problems.append(f"baris {n}: platform {source!r} nggak ada di daftar")
        if status not in STATUSES:
            problems.append(f"baris {n}: status {status!r} nggak ada di daftar")
        if not hasattr(applied, "strftime"):
            problems.append(f"baris {n}: tanggal lamar {applied!r} bukan tanggal")
            continue
        resp, req = details.pop((company, position), ([], []))
        rows.append(
            [company, position, applied.strftime("%Y-%m-%d"), text(url), text(location), source, status,
             text(notes), "\n".join(resp) or None, "\n".join(req) or None]
        )

    problems += [f"detail tanpa lamaran yang cocok: {k}" for k in details]
    if problems:
        sys.exit("Import dibatalkan:\n- " + "\n- ".join(problems))

    values = ",\n".join(
        "  (" + ", ".join(f"{sql(v)}::date" if i == 2 else sql(v) for i, v in enumerate(r)) + ")" for r in rows
    )
    print(f"""-- Dibuat oleh scripts/import_xlsx.py dari {len(rows)} lamaran. Jangan diedit manual.
insert into applications
  (user_id, company, position, applied_at, job_url, location, source, status, notes, responsibilities, requirements)
select u.id, v.*
from auth.users u,
(values
{values}
) as v(company, position, applied_at, job_url, location, source, status, notes, responsibilities, requirements)
where u.email = {sql(email)}
  and not exists (
    select 1 from applications a
    where a.user_id = u.id and a.company = v.company and a.position = v.position and a.applied_at = v.applied_at
  );""")
    print(f"-- {len(rows)} lamaran, {sum(1 for r in rows if r[8] or r[9])} dengan detail.", file=sys.stderr)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
